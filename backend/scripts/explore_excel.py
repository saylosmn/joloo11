import openpyxl
import os

path = "/app/backend/data/ZHD_800.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
print("SHEETS:", wb.sheetnames)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"\n===== SHEET: {sn} | dims={ws.dimensions} | max_row={ws.max_row} max_col={ws.max_column} =====")
    # print first 8 rows
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=8, values_only=True), start=1):
        # truncate long cells
        cells = []
        for c in row:
            s = str(c) if c is not None else ""
            if len(s) > 40:
                s = s[:40] + "…"
            cells.append(s)
        print(f"row{i}:", cells)

# Check for images
print("\n===== IMAGES CHECK =====")
try:
    wb2 = openpyxl.load_workbook(path)
    for sn in wb2.sheetnames:
        ws = wb2[sn]
        imgs = getattr(ws, "_images", [])
        print(f"Sheet {sn}: {len(imgs)} images")
        for im in imgs[:5]:
            anchor = im.anchor
            try:
                print("   anchor:", anchor._from.row, anchor._from.col)
            except Exception as e:
                print("   anchor(unknown):", e)
except Exception as e:
    print("img check err:", e)
