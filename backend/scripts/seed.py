"""
Seed script for the Mongolian Traffic Rules Exam app.

Reads the Excel file, runs validation, writes broken rows to skipped_rows.json,
then seeds `categories` and `questions` collections into MongoDB.

Run:  cd /app/backend && python scripts/seed.py
"""
import os
import json
import openpyxl
from pathlib import Path
from datetime import datetime, timezone
from pymongo import MongoClient
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parent.parent
load_dotenv(ROOT / ".env")

XLSX = str(ROOT / "data" / "ZHD_800.xlsx")
IMAGE_MAP = str(ROOT / "data" / "image_map.json")
SKIPPED = str(ROOT / "scripts" / "skipped_rows.json")

LETTERS = ["А", "Б", "В", "Г"]


def clean(x):
    return str(x).strip() if x is not None else ""


def main():
    # ---- image map (question № -> filename) ----
    with open(IMAGE_MAP, "r", encoding="utf-8") as f:
        raw = json.load(f)
    img_map = {int(k): v for k, v in raw.items()}

    wb = openpyxl.load_workbook(XLSX, data_only=True)
    sheet = "Асуулт" if "Асуулт" in wb.sheetnames else "Асуултууд"
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=2, values_only=True))

    broken = []
    seen = {}
    valid = []
    cat_order = []  # keep first-appearance order

    for idx, r in enumerate(rows):
        excel_row = idx + 2
        num = clean(r[0])
        cat = clean(r[1])
        q = clean(r[2])
        opts = [clean(r[4]), clean(r[5]), clean(r[6]), clean(r[7])]
        correct = clean(r[8]).upper()
        expl = clean(r[9])

        non_empty = [o for o in opts if o]
        has_image = num.isdigit() and int(num) in img_map
        reasons = []

        if not q:
            reasons.append("хоосон асуулт")
        if len(non_empty) < 2:
            reasons.append(f"хариулт 2-оос цөөн ({len(non_empty)})")
        if not cat:
            reasons.append("бүлэг тодорхойгүй")
        if correct not in LETTERS:
            reasons.append(f"зөв хариулт тодорхойгүй эсвэл олон: '{correct}'")
        else:
            ci = LETTERS.index(correct)
            if ci >= len(opts) or not opts[ci]:
                reasons.append(f"зөв хариултын багана хоосон: '{correct}'")
        if q and not has_image:
            key = q.lower() + "||" + "|".join(o.lower() for o in opts)
            if key in seen:
                reasons.append(f"давхардсан асуулт (мөр {seen[key]})")
            else:
                seen[key] = excel_row

        if reasons:
            broken.append({
                "excel_row": excel_row, "num": num,
                "reasons": reasons, "question": q[:80],
            })
            continue

        if cat not in cat_order:
            cat_order.append(cat)

        options = []
        for i, letter in enumerate(LETTERS):
            if opts[i]:
                options.append({
                    "key": letter,
                    "text": opts[i],
                    "isCorrect": letter == correct,
                })

        image_url = f"/api/images/{img_map[int(num)]}" if has_image else None
        valid.append({
            "num": int(num) if num.isdigit() else num,
            "category_name": cat,
            "question_text": q,
            "options": options,
            "correct_key": correct,
            "explanation": expl,
            "image_url": image_url,
            "source_row": excel_row,
        })

    # ---- write skipped report ----
    with open(SKIPPED, "w", encoding="utf-8") as f:
        json.dump(broken, f, ensure_ascii=False, indent=2)

    # ---- seed to Mongo ----
    mongo_url = os.environ["MONGO_URL"]
    db_name = os.environ["DB_NAME"]
    client = MongoClient(mongo_url)
    db = client[db_name]

    db.categories.delete_many({})
    db.questions.delete_many({})

    cat_docs = []
    cat_id_by_name = {}
    counts = {}
    for c in valid:
        counts[c["category_name"]] = counts.get(c["category_name"], 0) + 1

    for order, name in enumerate(cat_order):
        cid = f"cat_{order + 1}"
        cat_id_by_name[name] = cid
        cat_docs.append({
            "category_id": cid,
            "name": name,
            "order": order,
            "questionCount": counts.get(name, 0),
        })
    if cat_docs:
        db.categories.insert_many(cat_docs)

    q_docs = []
    now = datetime.now(timezone.utc).isoformat()
    for c in valid:
        q_docs.append({
            "question_id": f"q_{c['source_row']}",
            "num": c["num"],
            "category_id": cat_id_by_name[c["category_name"]],
            "category_name": c["category_name"],
            "questionText": c["question_text"],
            "options": c["options"],
            "correctKey": c["correct_key"],
            "explanation": c["explanation"],
            "imageUrl": c["image_url"],
            "sourceRow": c["source_row"],
            "createdAt": now,
        })
    if q_docs:
        db.questions.insert_many(q_docs)

    # ---- report ----
    print("=" * 60)
    print("SEED ТАЙЛАН")
    print("=" * 60)
    print(f"Excel-ийн нийт мөр (толгойгүй): {len(rows)}")
    print(f"Хүчинтэй асуулт (seed хийсэн):  {len(valid)}")
    print(f"Алдаатай/алгассан мөр:          {len(broken)}")
    print(f"Зурагтай асуулт:                {sum(1 for c in valid if c['image_url'])}")
    print(f"Нийт бүлэг:                     {len(cat_docs)}")
    print("-" * 60)
    print("БҮЛЭГ БҮРИЙН АСУУЛТЫН ТОО:")
    for cd in cat_docs:
        print(f"  [{cd['order']+1:>2}] {cd['name']}: {cd['questionCount']}")
    print("-" * 60)

    db_count = db.questions.count_documents({})
    print(f"MongoDB дахь асуултын тоо: {db_count}")
    if db_count == len(valid):
        print("✅ MongoDB-ийн тоо seed хийсэн хүчинтэй асуултын тоотой ТААРЧ БАЙНА.")
    else:
        print("❌ ЗӨРҮҮ ИЛЭРЛЭЭ!")
    if broken:
        print(f"⚠️  {len(broken)} мөр алгаслаа. Дэлгэрэнгүйг scripts/skipped_rows.json-оос үзнэ үү.")
    else:
        print("✅ Алдаатай мөр байхгүй.")
    client.close()


if __name__ == "__main__":
    main()
