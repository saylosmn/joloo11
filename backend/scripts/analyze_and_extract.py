"""
Analyze the Excel, print category summary, detect broken rows, and extract images.
Images are anchored to column D (Зураг). Save each image mapped to its question №.
"""
import openpyxl
import os
import json

PATH = "/app/backend/data/ZHD_800.xlsx"
IMG_DIR = "/app/backend/data/images"
os.makedirs(IMG_DIR, exist_ok=True)

wb = openpyxl.load_workbook(PATH)
ws = wb["Асуулт"] if "Асуулт" in wb.sheetnames else wb["Асуултууд"]

# ---- Extract images, map by anchor row -> question № ----
# anchor._from.row is 0-indexed. Excel row = anchor_row + 1. Question № = excel_row - 1.
img_map = {}  # qnum -> filename
count = 0
for im in getattr(ws, "_images", []):
    try:
        arow = im.anchor._from.row  # 0-indexed
        excel_row = arow + 1
        qnum = excel_row - 1
        ext = "png"
        try:
            ext = im.format.lower()
        except Exception:
            ext = "png"
        fname = f"q_{qnum}.{ext}"
        data = im._data()
        with open(os.path.join(IMG_DIR, fname), "wb") as f:
            f.write(data)
        img_map[qnum] = fname
        count += 1
    except Exception as e:
        print("img err", e)

print(f"Extracted {count} images. Mapped question numbers range: "
      f"{min(img_map) if img_map else None}..{max(img_map) if img_map else None}")
with open("/app/backend/data/image_map.json", "w", encoding="utf-8") as f:
    json.dump(img_map, f, ensure_ascii=False, indent=2)

# ---- Read all question rows (data_only for values) ----
wb2 = openpyxl.load_workbook(PATH, data_only=True)
ws2 = wb2["Асуулт"] if "Асуулт" in wb2.sheetnames else wb2["Асуултууд"]

rows = list(ws2.iter_rows(min_row=2, values_only=True))
print(f"\nTotal data rows: {len(rows)}")

# Column indices: 0 №, 1 Бүлэг, 2 Асуулт, 3 Зураг, 4 А, 5 Б, 6 В, 7 Г, 8 Зөв, 9 Тайлбар
cats = {}
broken = []
seen_text = {}
letters = ["А", "Б", "В", "Г"]

def clean(x):
    return str(x).strip() if x is not None else ""

for idx, r in enumerate(rows):
    excel_row = idx + 2
    num = clean(r[0])
    cat = clean(r[1])
    q = clean(r[2])
    opts = [clean(r[4]), clean(r[5]), clean(r[6]), clean(r[7])]
    correct = clean(r[8]).upper()
    expl = clean(r[9])

    non_empty_opts = [o for o in opts if o]
    reasons = []
    if not q:
        reasons.append("хоосон асуулт")
    if len(non_empty_opts) < 2:
        reasons.append(f"хариулт < 2 ({len(non_empty_opts)})")
    if not cat:
        reasons.append("бүлэг тодорхойгүй")
    if correct not in letters:
        reasons.append(f"зөв хариулт буруу/олон: '{correct}'")
    else:
        # check that the correct-letter option exists
        ci = letters.index(correct)
        if ci >= len(opts) or not opts[ci]:
            reasons.append(f"зөв хариултын багана хоосон: {correct}")
    has_image = num.isdigit() and int(num) in img_map
    if q and not has_image:
        # For image questions the text legitimately repeats (road signs),
        # so only dedup by text+options for non-image rows.
        key = q.lower() + "||" + "|".join(o.lower() for o in opts)
        if key in seen_text:
            reasons.append(f"давхардсан асуулт (мөр {seen_text[key]})")
        else:
            seen_text[key] = excel_row

    if reasons:
        broken.append({"excel_row": excel_row, "num": num, "reasons": reasons, "question": q[:60]})
    else:
        cats.setdefault(cat, 0)
        cats[cat] += 1

print("\n===== CATEGORY SUMMARY (valid rows) =====")
total_valid = 0
for c in sorted(cats.keys()):
    print(f"  {c}: {cats[c]}")
    total_valid += cats[c]
print(f"TOTAL VALID: {total_valid}")
print(f"BROKEN/SKIPPED: {len(broken)}")

with open("/app/backend/scripts/skipped_rows.json", "w", encoding="utf-8") as f:
    json.dump(broken, f, ensure_ascii=False, indent=2)

if broken:
    print("\n===== BROKEN ROWS (first 30) =====")
    for b in broken[:30]:
        print(f"  row {b['excel_row']} №{b['num']}: {b['reasons']}")
